"""
Serviço de geração de relatórios em PDF, Excel e CSV.
"""

import csv
import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import cm
from reportlab.platypus import (
    SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle


def generate_pdf_report(title: str, headers: list, rows: list, summary: dict = None) -> io.BytesIO:
    """Gera um relatório em PDF a partir de cabeçalhos e linhas de dados."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        topMargin=1.5 * cm,
        bottomMargin=1.5 * cm,
        leftMargin=1.5 * cm,
        rightMargin=1.5 * cm,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'CustomTitle', parent=styles['Heading1'],
        fontSize=18, textColor=colors.HexColor('#0d3b66'), spaceAfter=6
    )
    subtitle_style = ParagraphStyle(
        'CustomSubtitle', parent=styles['Normal'],
        fontSize=10, textColor=colors.HexColor('#555555'), spaceAfter=12
    )

    elements = [
        Paragraph(title, title_style),
        Paragraph(f'Gerado em {date.today().strftime("%d/%m/%Y")}', subtitle_style),
        Spacer(1, 0.4 * cm),
    ]

    table_data = [headers] + rows
    table = Table(table_data, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0d3b66')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cccccc')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f5f7fa')]),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))
    elements.append(table)

    if summary:
        elements.append(Spacer(1, 0.6 * cm))
        summary_style = ParagraphStyle(
            'Summary', parent=styles['Normal'], fontSize=10, spaceAfter=4
        )
        for key, value in summary.items():
            elements.append(Paragraph(f'<b>{key}:</b> {value}', summary_style))

    doc.build(elements)
    buffer.seek(0)
    return buffer


def generate_excel_report(title: str, headers: list, rows: list) -> io.BytesIO:
    """Gera um relatório em Excel (.xlsx)."""
    buffer = io.BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31] if title else 'Relatório'

    header_font = Font(bold=True, color='FFFFFF', size=11)
    header_fill = PatternFill(start_color='0D3B66', end_color='0D3B66', fill_type='solid')
    header_alignment = Alignment(horizontal='center', vertical='center')

    ws.append(headers)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    for row in rows:
        ws.append(row)

    for col_idx, header in enumerate(headers, start=1):
        max_length = len(str(header))
        for row in rows:
            value = row[col_idx - 1] if col_idx - 1 < len(row) else ''
            max_length = max(max_length, len(str(value)))
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = max_length + 4

    wb.save(buffer)
    buffer.seek(0)
    return buffer


def generate_csv_report(headers: list, rows: list) -> io.StringIO:
    """Gera um relatório em CSV."""
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=';')
    writer.writerow(headers)
    writer.writerows(rows)
    buffer.seek(0)
    return buffer
